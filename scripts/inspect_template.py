#!/usr/bin/env python3
"""Inspect PowerPoint template layouts and placeholders."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER


EMU_PER_INCH = 914400
ROLE_SCHEMA_VERSION = 1

CONTENT_ROLES = {
    "title",
    "body_content",
    "subtitle",
    "section_label",
    "caption",
    "source",
    "attribution",
    "quote",
    "callout",
    "footer_or_admin",
    "non_text",
    "unknown",
}
CONTENT_CLASSES = {
    "core_content",
    "optional_supporting",
    "administrative",
    "non_text",
    "review",
}

TITLE_TYPES = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE,
}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.VERTICAL_BODY}
MULTIPURPOSE_TYPES = {
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
    PP_PLACEHOLDER.MIXED,
}
ADMIN_TYPES = {
    PP_PLACEHOLDER.HEADER,
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.SLIDE_NUMBER,
}
NON_TEXT_TYPES = {
    PP_PLACEHOLDER.BITMAP,
    PP_PLACEHOLDER.CHART,
    PP_PLACEHOLDER.MEDIA_CLIP,
    PP_PLACEHOLDER.ORG_CHART,
    PP_PLACEHOLDER.PICTURE,
    PP_PLACEHOLDER.SLIDE_IMAGE,
    PP_PLACEHOLDER.TABLE,
}

# Specific supporting roles precede generic title/body names. Name-derived roles
# are always reviewable because corporate templates often use misleading names.
NAME_ROLE_RULES = (
    (r"\b(?:subtitle|subhead|subheading)\b", "subtitle", "optional_supporting"),
    (
        r"\b(?:eyebrow|kicker|overline|section label|category label)\b",
        "section_label",
        "optional_supporting",
    ),
    (r"\bcaption\b", "caption", "optional_supporting"),
    (
        r"\b(?:source|citation|reference|credit|evidence note)\b",
        "source",
        "optional_supporting",
    ),
    (
        r"\b(?:attribution|byline|author|speaker)\b",
        "attribution",
        "optional_supporting",
    ),
    (r"\b(?:quote|quotation|testimonial)\b", "quote", "optional_supporting"),
    (
        r"\b(?:takeaway|key message|insight|callout|implication|recommendation|highlight)\b",
        "callout",
        "optional_supporting",
    ),
    (
        r"\b(?:footer|header|confidential|copyright|legal)\b",
        "footer_or_admin",
        "administrative",
    ),
)


def emu_to_inches(value: int | None) -> float | None:
    if value is None:
        return None
    return round(value / EMU_PER_INCH, 3)


def _role_result(
    role: str,
    content_class: str,
    basis: str,
    confidence: str,
    review_required: bool,
    is_text_placeholder: bool,
) -> dict[str, Any]:
    return {
        "is_text_placeholder": is_text_placeholder,
        "content_role": role,
        "content_class": content_class,
        "role_basis": basis,
        "role_confidence": confidence,
        "review_required": review_required,
        "role_note": None,
    }


def normalize_placeholder_name(name: str) -> str:
    """Normalize a placeholder name for conservative word-boundary matching."""

    return " ".join(re.findall(r"[a-z0-9]+", name.casefold()))


def classify_placeholder(shape: Any) -> dict[str, Any]:
    """Infer a likely role without treating any placeholder as required content."""

    placeholder_type = shape.placeholder_format.type

    if placeholder_type in TITLE_TYPES:
        return _role_result(
            "title", "core_content", "placeholder_type", "high", False, True
        )
    if placeholder_type == PP_PLACEHOLDER.SUBTITLE:
        return _role_result(
            "subtitle",
            "optional_supporting",
            "placeholder_type",
            "high",
            False,
            True,
        )
    if placeholder_type in ADMIN_TYPES:
        return _role_result(
            "footer_or_admin",
            "administrative",
            "placeholder_type",
            "high",
            False,
            True,
        )
    if placeholder_type in NON_TEXT_TYPES:
        return _role_result(
            "non_text", "non_text", "placeholder_type", "high", False, False
        )

    has_text_frame = bool(getattr(shape, "has_text_frame", False))
    normalized_name = normalize_placeholder_name(str(getattr(shape, "name", "")))
    if has_text_frame:
        for pattern, role, content_class in NAME_ROLE_RULES:
            if re.search(pattern, normalized_name):
                return _role_result(
                    role,
                    content_class,
                    "placeholder_name",
                    "medium",
                    True,
                    True,
                )

    if placeholder_type in BODY_TYPES:
        return _role_result(
            "body_content",
            "core_content",
            "placeholder_type",
            "high",
            False,
            True,
        )
    if placeholder_type in MULTIPURPOSE_TYPES and has_text_frame:
        return _role_result(
            "body_content",
            "review",
            "ambiguous_placeholder_type",
            "low",
            True,
            True,
        )
    if has_text_frame:
        return _role_result(
            "unknown", "review", "unrecognized", "low", True, True
        )
    return _role_result(
        "non_text", "non_text", "no_text_frame", "low", True, False
    )


def _reject_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Duplicate JSON key: {key}")
        result[key] = value
    return result


def load_role_overrides(path: Path) -> dict[tuple[int, int], dict[str, Any]]:
    """Load and strictly validate template-specific placeholder role overrides."""

    try:
        with path.open("r", encoding="utf-8") as handle:
            data = json.load(handle, object_pairs_hook=_reject_duplicate_keys)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid role override JSON in {path}: {exc}") from exc

    if not isinstance(data, dict):
        raise ValueError("Role override file must contain a JSON object")
    unknown_top_level = set(data) - {"version", "overrides"}
    if unknown_top_level:
        raise ValueError(
            "Unknown role override top-level field(s): "
            + ", ".join(sorted(unknown_top_level))
        )
    if data.get("version") != ROLE_SCHEMA_VERSION:
        raise ValueError(
            f"Unsupported role override version: {data.get('version')!r}; "
            f"expected {ROLE_SCHEMA_VERSION}"
        )
    raw_overrides = data.get("overrides")
    if not isinstance(raw_overrides, dict):
        raise ValueError("Role override 'overrides' field must be a JSON object")

    allowed_fields = {
        "layout_name",
        "placeholder_name",
        "content_role",
        "content_class",
        "note",
    }
    overrides: dict[tuple[int, int], dict[str, Any]] = {}
    for key, value in raw_overrides.items():
        match = re.fullmatch(r"(\d+):(\d+)", key)
        if match is None:
            raise ValueError(
                f"Invalid role override key {key!r}; expected layout_index:placeholder_idx"
            )
        if not isinstance(value, dict):
            raise ValueError(f"Role override {key!r} must be a JSON object")
        unknown_fields = set(value) - allowed_fields
        if unknown_fields:
            raise ValueError(
                f"Unknown field(s) in role override {key!r}: "
                + ", ".join(sorted(unknown_fields))
            )

        role = value.get("content_role")
        content_class = value.get("content_class")
        if role not in CONTENT_ROLES:
            raise ValueError(f"Unknown content_role in role override {key!r}: {role!r}")
        if content_class not in CONTENT_CLASSES:
            raise ValueError(
                f"Unknown content_class in role override {key!r}: {content_class!r}"
            )
        for string_field in ("layout_name", "placeholder_name", "note"):
            if string_field in value and not isinstance(value[string_field], str):
                raise ValueError(
                    f"Role override {key!r} field {string_field!r} must be a string"
                )

        target = (int(match.group(1)), int(match.group(2)))
        overrides[target] = dict(value)
    return overrides


def apply_role_override(
    inferred: dict[str, Any],
    override: dict[str, Any],
    *,
    layout_index: int,
    layout_name: str,
    placeholder_name: str,
    placeholder_idx: int,
) -> dict[str, Any]:
    """Apply a reviewed override after checking optional template-drift guards."""

    target = f"{layout_index}:{placeholder_idx}"
    guarded_layout_name = override.get("layout_name")
    if guarded_layout_name is not None and guarded_layout_name != layout_name:
        raise ValueError(
            f"Role override {target} expected layout {guarded_layout_name!r}, "
            f"found {layout_name!r}"
        )
    guarded_placeholder_name = override.get("placeholder_name")
    if (
        guarded_placeholder_name is not None
        and guarded_placeholder_name != placeholder_name
    ):
        raise ValueError(
            f"Role override {target} expected placeholder {guarded_placeholder_name!r}, "
            f"found {placeholder_name!r}"
        )

    result = dict(inferred)
    result.update(
        {
            "content_role": override["content_role"],
            "content_class": override["content_class"],
            "role_basis": "override",
            "role_confidence": "explicit",
            "review_required": False,
            "role_note": override.get("note"),
        }
    )
    return result


def placeholder_info(
    shape: Any,
    *,
    layout_index: int,
    layout_name: str,
    override: dict[str, Any] | None = None,
) -> dict[str, Any]:
    fmt = shape.placeholder_format
    role_info = classify_placeholder(shape)
    if override is not None:
        role_info = apply_role_override(
            role_info,
            override,
            layout_index=layout_index,
            layout_name=layout_name,
            placeholder_name=shape.name,
            placeholder_idx=fmt.idx,
        )
    return {
        "idx": fmt.idx,
        "type": str(fmt.type),
        "name": shape.name,
        "left_in": emu_to_inches(shape.left),
        "top_in": emu_to_inches(shape.top),
        "width_in": emu_to_inches(shape.width),
        "height_in": emu_to_inches(shape.height),
        **role_info,
    }


def inspect(
    path: Path,
    role_overrides: dict[tuple[int, int], dict[str, Any]] | None = None,
) -> dict[str, Any]:
    prs = Presentation(str(path))
    overrides = role_overrides or {}
    unmatched_override_targets = set(overrides)
    layouts = []
    for index, layout in enumerate(prs.slide_layouts):
        placeholders = []
        for shape in layout.placeholders:
            target = (index, shape.placeholder_format.idx)
            override = overrides.get(target)
            if override is not None:
                unmatched_override_targets.discard(target)
            placeholders.append(
                placeholder_info(
                    shape,
                    layout_index=index,
                    layout_name=layout.name,
                    override=override,
                )
            )
        layouts.append(
            {
                "index": index,
                "name": layout.name,
                "placeholder_count": len(placeholders),
                "review_required_count": sum(
                    placeholder["review_required"] for placeholder in placeholders
                ),
                "placeholders": placeholders,
            }
        )
    if unmatched_override_targets:
        targets = ", ".join(
            f"{layout_index}:{placeholder_idx}"
            for layout_index, placeholder_idx in sorted(unmatched_override_targets)
        )
        raise ValueError(f"Role override target(s) not found in template: {targets}")
    return {
        "path": str(path),
        "slide_width_in": emu_to_inches(prs.slide_width),
        "slide_height_in": emu_to_inches(prs.slide_height),
        "layout_count": len(layouts),
        "role_override_count": len(overrides),
        "review_required_count": sum(
            layout["review_required_count"] for layout in layouts
        ),
        "layouts": layouts,
    }


def markdown_cell(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def to_markdown(report: dict[str, Any]) -> str:
    lines = [
        f"# Template Audit: {report['path']}",
        "",
        f"- Size: {report['slide_width_in']} x {report['slide_height_in']} in",
        f"- Layouts: {report['layout_count']}",
        f"- Role overrides: {report.get('role_override_count', 0)}",
        f"- Placeholders requiring role review: {report.get('review_required_count', 0)}",
        "",
    ]
    for layout in report["layouts"]:
        lines.append(f"## {layout['index']}: {layout['name']}")
        if not layout["placeholders"]:
            lines.append("")
            lines.append("_No placeholders._")
            lines.append("")
            continue
        lines.append("")
        lines.append(
            "| idx | type | name | role | class | basis | review | left | top | width | height |"
        )
        lines.append(
            "| ---: | --- | --- | --- | --- | --- | --- | ---: | ---: | ---: | ---: |"
        )
        for ph in layout["placeholders"]:
            lines.append(
                "| {idx} | {type} | {name} | {role} | {content_class} | {basis} | {review} | {left_in} | {top_in} | {width_in} | {height_in} |".format(
                    idx=ph["idx"],
                    type=markdown_cell(ph["type"]),
                    name=markdown_cell(ph["name"]),
                    role=ph["content_role"],
                    content_class=ph["content_class"],
                    basis=ph["role_basis"],
                    review="yes" if ph["review_required"] else "no",
                    left_in=ph["left_in"],
                    top_in=ph["top_in"],
                    width_in=ph["width_in"],
                    height_in=ph["height_in"],
                )
            )
        lines.append("")
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path, help="Path to .pptx template/deck")
    parser.add_argument(
        "--format",
        choices=("json", "markdown"),
        default="json",
        help="Output format",
    )
    parser.add_argument(
        "--role-overrides",
        type=Path,
        help="Path to reviewed JSON placeholder-role overrides",
    )
    args = parser.parse_args()

    try:
        role_overrides = (
            load_role_overrides(args.role_overrides) if args.role_overrides else {}
        )
        report = inspect(args.template, role_overrides)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    if args.format == "markdown":
        print(to_markdown(report))
    else:
        print(json.dumps(report, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
