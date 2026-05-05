#!/usr/bin/env python3
"""Create a small deck containing an editable native PowerPoint chart."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches


def build_chart_data() -> CategoryChartData:
    chart_data = CategoryChartData(number_format="#,##0.0")
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("Actual", (12.4, 14.2, 13.8, 16.1))
    chart_data.add_series("Target", (11.5, 13.0, 14.0, 15.0))
    return chart_data


def find_layout(prs: Presentation):
    for layout in prs.slide_layouts:
        types = {shape.placeholder_format.type for shape in layout.placeholders}
        if PP_PLACEHOLDER.CHART in types and PP_PLACEHOLDER.TITLE in types:
            return layout
    for layout in prs.slide_layouts:
        if any(shape.placeholder_format.type == PP_PLACEHOLDER.TITLE for shape in layout.placeholders):
            return layout
    return prs.slide_layouts[0]


def find_placeholder(slide, placeholder_type):
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.type == placeholder_type:
            return placeholder
    return None


def add_chart(slide, chart_data):
    chart_placeholder = find_placeholder(slide, PP_PLACEHOLDER.CHART)
    if chart_placeholder is not None:
        return chart_placeholder.insert_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, chart_data)

    content_placeholder = find_placeholder(slide, PP_PLACEHOLDER.OBJECT)
    if content_placeholder is not None:
        return slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            content_placeholder.left,
            content_placeholder.top,
            content_placeholder.width,
            content_placeholder.height,
            chart_data,
        )

    return slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1.0),
        Inches(1.6),
        Inches(8.0),
        Inches(4.5),
        chart_data,
    )


def create_deck(template: Path | None, output: Path) -> None:
    prs = Presentation(str(template)) if template else Presentation()
    slide = prs.slides.add_slide(find_layout(prs))

    if slide.shapes.title:
        slide.shapes.title.text = "Native editable chart"

    chart = add_chart(slide, build_chart_data()).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.chart_style = 10

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, help="Optional .pptx template/deck")
    parser.add_argument("--output", type=Path, required=True, help="Output .pptx path")
    args = parser.parse_args()

    create_deck(args.template, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
