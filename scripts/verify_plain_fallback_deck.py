#!/usr/bin/env python3
"""Verify the generated plain fallback PowerPoint deck.

The checks enforce the deck's purpose: example slides must survive a paste
into a corporate template with *Use Destination Theme* and come out on-brand.
That requires:

* one master with title, body, date, footer and slide-number placeholders
  sized for 16:9;
* the expected layouts, in order, with PowerPoint's standard names and types
  for the core layouts;
* layouts and the master that contain placeholders only (no decoration);
* title placeholders without ``idx="0"`` (it breaks inheritance outside
  PowerPoint);
* no hard-coded RGB colours or typeface names anywhere below the theme;
* placeholder text on slides without size, colour or font overrides;
* no embedded pictures;
* native, editable charts with restrained formatting and correct data.

It inspects the OOXML package directly, then loads the deck with python-pptx
as a final consumer-level check.
"""

from __future__ import annotations

import argparse
import json
import posixpath
import re
import zipfile
from pathlib import Path
from typing import Any, Iterable

from lxml import etree
from pptx import Presentation

try:  # Direct script execution places this directory first on sys.path.
    from build_plain_fallback_deck import CATALOG, LAYOUTS, STANDARD_LAYOUT_TYPES
except ModuleNotFoundError:  # Importing as scripts.verify_plain_fallback_deck.
    from scripts.build_plain_fallback_deck import CATALOG, LAYOUTS, STANDARD_LAYOUT_TYPES


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

NS = {"p": P_NS, "a": A_NS, "c": C_NS, "r": R_NS}
REL = {"rel": REL_NS}
CT = {"ct": CT_NS}

SLIDE_MASTER_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster"
SLIDE_LAYOUT_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
SLIDE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
CHART_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
PACKAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"

SLIDE_LAYOUT_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"
CHART_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
WORKBOOK_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DEFAULT_DECK = Path("references/plain-fallback-layouts.pptx")

EXPECTED_LAYOUTS = tuple((spec.name, spec.ooxml_type) for spec in LAYOUTS)
EXPECTED_SLIDE_COUNT = len(CATALOG)
ENTRY_BY_SLIDE = {entry.number: entry for entry in CATALOG}
NATIVE_CHART_SLIDES = frozenset(e.number for e in CATALOG if e.kind == "chart")
TABLE_SLIDES = frozenset(e.number for e in CATALOG if e.kind in ("table", "heatmap"))
DIAGRAM_SLIDES = frozenset(e.number for e in CATALOG if e.kind in ("diagram", "treemap"))
DESCENDING_BAR_SLIDES = frozenset({54, 55, 56, 57, 58, 67})
WATERFALL_SLIDE = 64
MIN_DIAGRAM_SHAPES = 4

# The standard Office layouts every corporate template is expected to have.
REQUIRED_STANDARD_LAYOUTS = {
    "Title Slide": "title",
    "Title and Content": "obj",
    "Section Header": "secHead",
    "Two Content": "twoObj",
    "Comparison": "twoTxTwoObj",
    "Title Only": "titleOnly",
    "Blank": "blank",
    "Content with Caption": "objTx",
    "Picture with Caption": "picTx",
}

EXPECTED_PLOTS: dict[int, tuple[tuple[str, str | None, str | None], ...]] = {
    49: (("barChart", "col", "clustered"),),
    50: (("barChart", "col", "clustered"),),
    51: (("barChart", "col", "clustered"),),
    52: (("barChart", "col", "stacked"),),
    53: (("barChart", "col", "percentStacked"),),
    54: (("barChart", "bar", "clustered"),),
    55: (("barChart", "bar", "clustered"),),
    56: (("barChart", "bar", "stacked"),),
    57: (("barChart", "bar", "percentStacked"),),
    58: (("barChart", "bar", "clustered"),),
    59: (("pieChart", None, None),),
    60: (("doughnutChart", None, None),),
    61: (("lineChart", None, "standard"),),
    62: (("areaChart", None, "stacked"),),
    63: (("barChart", "col", "clustered"), ("lineChart", None, "standard")),
    64: (("barChart", "col", "stacked"),),
    65: (("scatterChart", None, None),),
    66: (("bubbleChart", None, None),),
    67: (("barChart", "bar", "clustered"),),
    72: (("barChart", "col", "clustered"),),
}

PROHIBITED_COPY_PATTERNS = (
    ("where-insight-meets-impact", re.compile(r"\bwhere\s+insight\s+meets\s+impact\b", re.I)),
    ("smarter-path-forward", re.compile(r"\ba\s+smarter\s+path\s+forward\b", re.I)),
    ("vision-to-value", re.compile(r"\bfrom\s+vision\s+to\s+value\b", re.I)),
    ("unlocking-generic-value", re.compile(r"\bunlock(?:ing)?\s+(?:the\s+)?(?:power|potential|value)\b", re.I)),
    ("where-x-meets-y", re.compile(r"\bwhere\s+[a-z][a-z -]{1,30}\s+meets\s+[a-z][a-z -]{1,30}\b", re.I)),
)

HARD_CODED_COLOR_TAGS = ("srgbClr", "scrgbClr", "sysClr", "hslClr", "prstClr")
FONT_TAGS = ("latin", "ea", "cs", "sym")
SHAPE_TAGS = ("sp", "pic", "graphicFrame", "cxnSp", "grpSp")


def expected_layout_names() -> tuple[str, ...]:
    return tuple(name for name, _ in EXPECTED_LAYOUTS)


def _rels_name(part_name: str) -> str:
    directory, filename = posixpath.split(part_name)
    return posixpath.join(directory, "_rels", filename + ".rels")


def _resolve_target(source_part: str, target: str) -> str:
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _is_physical_part(name: str, directory: str, stem: str) -> bool:
    return bool(re.fullmatch(rf"{re.escape(directory)}/{re.escape(stem)}\d+\.xml", name))


def _relationship_map(root: etree._Element | None) -> dict[str, etree._Element]:
    if root is None:
        return {}
    return {rel.get("Id", ""): rel for rel in root.findall("rel:Relationship", REL)}


def _placeholder_type(ph: etree._Element) -> str:
    return ph.get("type", "obj")  # OOXML default when @type is absent


def _part_content_type(part_name: str, overrides: dict[str, str], defaults: dict[str, str]) -> str | None:
    if part_name in overrides:
        return overrides[part_name]
    suffix = posixpath.basename(part_name).rsplit(".", 1)
    return defaults.get(suffix[1].lower() if len(suffix) == 2 else "")


def _parse_xml(parts: dict[str, bytes], part_name: str, errors: list[str]) -> etree._Element | None:
    payload = parts.get(part_name)
    if payload is None:
        errors.append(f"Missing required part: {part_name}")
        return None
    try:
        return etree.fromstring(payload)
    except (etree.XMLSyntaxError, ValueError) as exc:
        errors.append(f"Invalid XML in {part_name}: {exc}")
        return None


def _relationship_target(source_part: str, rel: etree._Element, parts: dict[str, bytes],
                         errors: list[str]) -> str | None:
    target = rel.get("Target")
    if not target:
        errors.append(f"Relationship {rel.get('Id')} in {source_part} has no target")
        return None
    if rel.get("TargetMode") == "External":
        errors.append(f"Relationship {rel.get('Id')} in {source_part} unexpectedly has an external target")
        return None
    resolved = _resolve_target(source_part, target)
    if resolved not in parts:
        errors.append(f"Relationship {rel.get('Id')} in {source_part} resolves to missing part {resolved}")
    return resolved


# ---------------------------------------------------------------------------
# Theme-inheritance checks
# ---------------------------------------------------------------------------


def _check_no_hard_coded_style(label: str, root: etree._Element, errors: list[str]) -> None:
    """Colours must be scheme colours and typefaces must be theme references."""

    for tag in HARD_CODED_COLOR_TAGS:
        found = root.findall(f".//a:{tag}", NS)
        if found:
            errors.append(f"{label} contains {len(found)} hard-coded colour(s) (a:{tag}); use theme colours")
    for tag in FONT_TAGS:
        for element in root.findall(f".//a:{tag}", NS):
            typeface = element.get("typeface", "")
            if typeface and not typeface.startswith("+"):
                errors.append(f"{label} names typeface {typeface!r}; use the theme fonts (+mj-lt/+mn-lt)")


def _check_placeholder_text_inherits(label: str, root: etree._Element, errors: list[str]) -> None:
    """Placeholder text on slides must inherit size, colour and font."""

    for shape in root.findall(".//p:sp", NS):
        if shape.find("p:nvSpPr/p:nvPr/p:ph", NS) is None:
            continue
        sp_pr = shape.find("p:spPr", NS)
        if sp_pr is not None and len(sp_pr):
            local_names = sorted({etree.QName(child).localname for child in sp_pr})
            errors.append(f"{label} placeholder overrides layout geometry or fill ({', '.join(local_names)})")
        for run_props in shape.findall(".//a:rPr", NS) + shape.findall(".//a:endParaRPr", NS):
            if run_props.get("sz"):
                errors.append(f"{label} placeholder text overrides font size; let the layout set it")
            if len(run_props):
                local_names = sorted({etree.QName(child).localname for child in run_props})
                errors.append(f"{label} placeholder text overrides {', '.join(local_names)}")


def _check_layout_structure(index: int, name: str, root: etree._Element, errors: list[str]) -> None:
    sp_tree = root.find("p:cSld/p:spTree", NS)
    if sp_tree is None:
        errors.append(f"Layout {index} ({name}) has no shape tree")
        return
    for child in sp_tree:
        local = etree.QName(child).localname
        if local not in SHAPE_TAGS:
            continue
        if child.find(".//p:nvPr/p:ph", NS) is None:
            shape_name = child.find(".//p:cNvPr", NS)
            errors.append(
                f"Layout {index} ({name}) contains a non-placeholder {local} "
                f"{shape_name.get('name') if shape_name is not None else ''!r}; layouts must hold placeholders only"
            )
    placeholders = root.findall(".//p:ph", NS)
    indices = [ph.get("idx", "0") for ph in placeholders]
    if len(indices) != len(set(indices)):
        errors.append(f"Layout {index} ({name}) contains duplicate placeholder indices")
    for ph in placeholders:
        if _placeholder_type(ph) in ("title", "ctrTitle") and ph.get("idx") is not None:
            errors.append(
                f"Layout {index} ({name}) title placeholder carries idx={ph.get('idx')!r}; "
                "omit idx on titles so every consumer inherits the layout position"
            )


def _check_master(master_part: str, master: etree._Element, slide_width: int, errors: list[str]) -> None:
    sp_tree = master.find("p:cSld/p:spTree", NS)
    types = []
    for child in sp_tree if sp_tree is not None else []:
        local = etree.QName(child).localname
        if local not in SHAPE_TAGS:
            continue
        ph = child.find(".//p:nvPr/p:ph", NS)
        if ph is None:
            errors.append(f"Slide master contains a non-placeholder {local}; keep brand furniture out of the fallback")
            continue
        types.append(_placeholder_type(ph))
        if _placeholder_type(ph) == "title" and ph.get("idx") is not None:
            errors.append("Slide master title placeholder carries an idx attribute")
        off = child.find("p:spPr/a:xfrm/a:off", NS)
        ext = child.find("p:spPr/a:xfrm/a:ext", NS)
        if _placeholder_type(ph) in ("title", "body") and off is not None and ext is not None:
            right_edge = int(off.get("x")) + int(ext.get("cx"))
            if right_edge < slide_width * 0.85:
                errors.append(
                    f"Slide master {_placeholder_type(ph)} placeholder ends at {right_edge} EMU on a "
                    f"{slide_width} EMU slide; size master placeholders for the slide width"
                )
    for required in ("title", "body", "dt", "ftr", "sldNum"):
        if required not in types:
            errors.append(f"Slide master lacks a {required} placeholder")
    _check_no_hard_coded_style("Slide master", master, errors)


# ---------------------------------------------------------------------------
# Chart checks
# ---------------------------------------------------------------------------


def _plot_signature(chart_root: etree._Element) -> list[tuple[str, str | None, str | None]]:
    actual = []
    for tag in ("barChart", "pieChart", "doughnutChart", "lineChart", "areaChart", "scatterChart", "bubbleChart"):
        for plot in chart_root.findall(f".//c:{tag}", NS):
            bar_dir = plot.find("c:barDir", NS)
            grouping = plot.find("c:grouping", NS)
            actual.append((tag, bar_dir.get("val") if bar_dir is not None else None,
                           grouping.get("val") if grouping is not None else None))
    return actual


def _series_values(series: etree._Element) -> list[float]:
    points = series.findall("c:val/c:numRef/c:numCache/c:pt", NS) or series.findall("c:val/c:numLit/c:pt", NS)
    indexed = []
    for point in points:
        value = point.find("c:v", NS)
        if value is not None and value.text is not None:
            indexed.append((int(point.get("idx", "0")), float(value.text)))
    return [value for _, value in sorted(indexed)]


def _check_descending_bar_order(slide: int, chart_root: etree._Element, errors: list[str]) -> None:
    orientations = [axis.find("c:scaling/c:orientation", NS) for axis in chart_root.findall(".//c:catAx", NS)]
    if len(orientations) != 1 or any(o is None or o.get("val") != "maxMin" for o in orientations):
        errors.append(
            f"Slide {slide} horizontal bar chart must reverse its category axis so the highest value appears at the top"
        )
    series_values = [_series_values(s) for s in chart_root.findall(".//c:barChart/c:ser", NS)]
    series_values = [values for values in series_values if values]
    if not series_values or len({len(v) for v in series_values}) != 1:
        errors.append(f"Slide {slide} horizontal bar chart has incomplete cached values")
        return
    totals = [round(sum(values), 6) for values in zip(*series_values)]
    if any(later > earlier + 1e-6 for earlier, later in zip(totals, totals[1:])):
        errors.append(f"Slide {slide} horizontal bar categories are not sorted highest-to-lowest: {totals!r}")


def _check_waterfall(slide: int, chart_root: etree._Element, errors: list[str]) -> None:
    series = chart_root.findall(".//c:barChart/c:ser", NS)
    if len(series) != 2:
        errors.append(f"Waterfall slide {slide} must have an invisible base series and a movement series")
        return
    base, movement = series
    fill = base.find("c:spPr/a:noFill", NS)
    if fill is None:
        errors.append(f"Waterfall slide {slide} base series must have no fill")
    base_labels = base.find("c:dLbls/c:showVal", NS)
    if base_labels is not None and base_labels.get("val") == "1":
        errors.append(f"Waterfall slide {slide} labels its invisible base series")
    movement_labels = movement.find("c:dLbls/c:showVal", NS)
    if movement_labels is None or movement_labels.get("val") != "1":
        errors.append(f"Waterfall slide {slide} must label its movement series")
    base_values, moves = _series_values(base), _series_values(movement)
    if len(base_values) < 3 or len(base_values) != len(moves):
        errors.append(f"Waterfall slide {slide} has incomplete cached values")
        return
    start, end, steps = moves[0], moves[-1], moves[1:-1]
    if abs(start + sum(steps) - end) > 1e-6:
        errors.append(f"Waterfall slide {slide} does not add up: {start} + {steps} != {end}")
    running = start
    for base_value, step in zip(base_values[1:-1], steps):
        if abs(base_value - running) > 1e-6:
            errors.append(f"Waterfall slide {slide} base {base_value} does not match running total {running}")
        running += step


def _check_chart_part(slide: int, chart_part: str, parts: dict[str, bytes], overrides: dict[str, str],
                      defaults: dict[str, str], errors: list[str]) -> str | None:
    chart_root = _parse_xml(parts, chart_part, errors)
    if chart_root is None:
        return None
    content_type = _part_content_type(chart_part, overrides, defaults)
    if content_type != CHART_CONTENT_TYPE:
        errors.append(f"Chart part {chart_part} has content type {content_type!r}, expected {CHART_CONTENT_TYPE!r}")
    if chart_root.findall(".//c:title", NS):
        errors.append(f"Slide {slide} chart {chart_part} contains a chart or axis title")
    if chart_root.findall(".//c:majorGridlines", NS):
        errors.append(f"Slide {slide} chart {chart_part} contains major gridlines")
    if chart_root.findall(".//c:minorGridlines", NS):
        errors.append(f"Slide {slide} chart {chart_part} contains minor gridlines")
    if chart_root.findall(".//a:blip", NS):
        errors.append(f"Slide {slide} chart {chart_part} contains an embedded image")
    _check_no_hard_coded_style(f"Slide {slide} chart {chart_part}", chart_root, errors)

    expected = EXPECTED_PLOTS.get(slide)
    actual = _plot_signature(chart_root)
    if expected is not None and sorted(actual) != sorted(expected):
        errors.append(f"Slide {slide} chart {chart_part} has plot contract {actual!r}; expected {list(expected)!r}")
    if slide in DESCENDING_BAR_SLIDES:
        _check_descending_bar_order(slide, chart_root, errors)
    if slide == WATERFALL_SLIDE:
        _check_waterfall(slide, chart_root, errors)
    if slide == 65:
        for series in chart_root.findall(".//c:scatterChart/c:ser", NS):
            if series.find("c:spPr/a:ln/a:noFill", NS) is None:
                errors.append(f"Slide {slide} scatter series must be markers only (no connecting line)")

    external = chart_root.findall(".//c:externalData", NS)
    if len(external) != 1:
        errors.append(f"Slide {slide} chart {chart_part} has {len(external)} externalData elements; expected one")
        return None
    rid = external[0].get(f"{{{R_NS}}}id")
    workbook_rel = _relationship_map(_parse_xml(parts, _rels_name(chart_part), errors)).get(rid or "")
    if workbook_rel is None:
        errors.append(f"Slide {slide} chart {chart_part} externalData relationship {rid!r} is missing")
        return None
    if workbook_rel.get("Type") != PACKAGE_REL:
        errors.append(f"Slide {slide} chart {chart_part} externalData relationship is not an embedded package")
    workbook_part = _relationship_target(chart_part, workbook_rel, parts, errors)
    if workbook_part is None:
        return None
    if _part_content_type(workbook_part, overrides, defaults) != WORKBOOK_CONTENT_TYPE:
        errors.append(f"Embedded workbook {workbook_part} lacks the workbook content type")
    return workbook_part


def _check_prohibited_copy(slide: int, slide_root: etree._Element, errors: list[str]) -> None:
    for text_element in slide_root.findall(".//a:t", NS):
        text = " ".join((text_element.text or "").split())
        for label, pattern in PROHIBITED_COPY_PATTERNS:
            if text and pattern.search(text):
                errors.append(f"Slide {slide} contains prohibited ornamental copy ({label}): {text!r}")


# ---------------------------------------------------------------------------
# Deck inspection
# ---------------------------------------------------------------------------


def inspect_deck(path: Path) -> dict[str, Any]:
    """Return a machine-readable acceptance report for *path*."""

    path = Path(path)
    errors: list[str] = []
    warnings: list[str] = []
    report: dict[str, Any] = {
        "path": str(path),
        "master_count": 0,
        "layout_count": 0,
        "standard_layout_count": 0,
        "slide_count": 0,
        "native_chart_part_count": 0,
        "embedded_workbook_count": 0,
        "table_slide_count": 0,
        "diagram_slide_count": 0,
        "errors": errors,
        "warnings": warnings,
        "ok": False,
    }
    if not path.is_file():
        errors.append(f"Deck not found: {path}")
        return report

    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            if len(names) != len(set(names)):
                errors.append("PPTX contains duplicate ZIP member names")
            corrupt = zf.testzip()
            if corrupt:
                errors.append(f"PPTX ZIP CRC failure in {corrupt}")
            parts = {name: zf.read(name) for name in names}
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        errors.append(f"Cannot read PPTX package: {exc}")
        return report

    content_types = _parse_xml(parts, "[Content_Types].xml", errors)
    overrides: dict[str, str] = {}
    defaults: dict[str, str] = {}
    if content_types is not None:
        overrides = {i.get("PartName", "").lstrip("/"): i.get("ContentType", "")
                     for i in content_types.findall("ct:Override", CT)}
        defaults = {i.get("Extension", "").lower(): i.get("ContentType", "")
                    for i in content_types.findall("ct:Default", CT)}

    media = sorted(name for name in parts if name.startswith("ppt/media/"))
    if media:
        errors.append(f"Deck embeds media {media!r}; leave picture placeholders empty instead of stock images")

    physical_masters = sorted(n for n in parts if _is_physical_part(n, "ppt/slideMasters", "slideMaster"))
    physical_layouts = sorted(n for n in parts if _is_physical_part(n, "ppt/slideLayouts", "slideLayout"))
    physical_slides = sorted(n for n in parts if _is_physical_part(n, "ppt/slides", "slide"))
    report["master_count"] = len(physical_masters)
    if len(physical_masters) != 1:
        errors.append(f"Expected exactly one physical slide master, found {len(physical_masters)}")
    if len(physical_layouts) != len(EXPECTED_LAYOUTS):
        errors.append(f"Expected exactly {len(EXPECTED_LAYOUTS)} physical slide layouts, found {len(physical_layouts)}")
    if len(physical_slides) != EXPECTED_SLIDE_COUNT:
        errors.append(f"Expected exactly {EXPECTED_SLIDE_COUNT} physical slides, found {len(physical_slides)}")

    presentation_part = "ppt/presentation.xml"
    presentation = _parse_xml(parts, presentation_part, errors)
    presentation_rel_map = _relationship_map(_parse_xml(parts, _rels_name(presentation_part), errors))
    if presentation is None:
        return report
    slide_size = presentation.find("p:sldSz", NS)
    slide_width = int(slide_size.get("cx")) if slide_size is not None else 12192000
    if slide_size is not None and abs(int(slide_size.get("cx")) / int(slide_size.get("cy")) - 16 / 9) > 0.01:
        errors.append("Slide size is not 16:9")

    master_ids = presentation.findall("p:sldMasterIdLst/p:sldMasterId", NS)
    if len(master_ids) != 1:
        errors.append(f"Expected presentation to reference one slide master, found {len(master_ids)}")
    master_part = None
    if master_ids:
        rid = master_ids[0].get(f"{{{R_NS}}}id")
        rel = presentation_rel_map.get(rid or "")
        if rel is None:
            errors.append(f"Presentation slide-master relationship {rid!r} is missing")
        else:
            master_part = _relationship_target(presentation_part, rel, parts, errors)

    layout_parts: list[str] = []
    layout_name_by_part: dict[str, str] = {}
    if master_part:
        master = _parse_xml(parts, master_part, errors)
        if master is not None:
            _check_master(master_part, master, slide_width, errors)
        master_rel_map = _relationship_map(_parse_xml(parts, _rels_name(master_part), errors))
        layout_ids = master.findall("p:sldLayoutIdLst/p:sldLayoutId", NS) if master is not None else []
        report["layout_count"] = len(layout_ids)
        if len(layout_ids) != len(EXPECTED_LAYOUTS):
            errors.append(f"Expected master to reference {len(EXPECTED_LAYOUTS)} layouts, found {len(layout_ids)}")
        seen_ids: set[str] = set()
        found_types: dict[str, str] = {}
        for index, layout_id in enumerate(layout_ids, start=1):
            numeric_id = layout_id.get("id", "")
            if numeric_id in seen_ids:
                errors.append(f"Duplicate master layout ID {numeric_id!r}")
            seen_ids.add(numeric_id)
            rid = layout_id.get(f"{{{R_NS}}}id")
            rel = master_rel_map.get(rid or "")
            if rel is None:
                errors.append(f"Master layout {index} relationship {rid!r} is missing")
                continue
            if rel.get("Type") != SLIDE_LAYOUT_REL:
                errors.append(f"Master relationship {rid!r} for layout {index} has wrong type")
            layout_part = _relationship_target(master_part, rel, parts, errors)
            if layout_part is None:
                continue
            layout_parts.append(layout_part)
            if _part_content_type(layout_part, overrides, defaults) != SLIDE_LAYOUT_CONTENT_TYPE:
                errors.append(f"Layout {index} part {layout_part} lacks the slide-layout content type")
            layout_root = _parse_xml(parts, layout_part, errors)
            if layout_root is None:
                continue
            c_sld = layout_root.find("p:cSld", NS)
            actual_name = c_sld.get("name") if c_sld is not None else None
            actual_type = layout_root.get("type", "cust")
            layout_name_by_part[layout_part] = actual_name or ""
            found_types[actual_name or ""] = actual_type
            if index <= len(EXPECTED_LAYOUTS):
                expected_name, expected_type = EXPECTED_LAYOUTS[index - 1]
                if actual_name != expected_name:
                    errors.append(f"Layout {index} is named {actual_name!r}; expected {expected_name!r}")
                if actual_type != expected_type:
                    errors.append(f"Layout {index} ({actual_name}) has type {actual_type!r}; expected {expected_type!r}")
            _check_layout_structure(index, actual_name or "", layout_root, errors)
            _check_no_hard_coded_style(f"Layout {index} ({actual_name})", layout_root, errors)
            back = [r for r in _parse_xml(parts, _rels_name(layout_part), errors).findall("rel:Relationship", REL)
                    if r.get("Type") == SLIDE_MASTER_REL] if _rels_name(layout_part) in parts else []
            if len(back) != 1:
                errors.append(f"Layout {index} must have exactly one back-relationship to the master; found {len(back)}")
            elif _relationship_target(layout_part, back[0], parts, errors) != master_part:
                errors.append(f"Layout {index} back-relationship does not target {master_part}")
        for name, ooxml_type in REQUIRED_STANDARD_LAYOUTS.items():
            if found_types.get(name) != ooxml_type:
                errors.append(
                    f"Standard layout {name!r} with type {ooxml_type!r} is missing; pasted slides will not map "
                    "onto the destination template"
                )
            else:
                report["standard_layout_count"] += 1
        if set(layout_parts) != set(physical_layouts):
            errors.append("Master/layout package mismatch")

    slide_ids = presentation.findall("p:sldIdLst/p:sldId", NS)
    report["slide_count"] = len(slide_ids)
    if len(slide_ids) != EXPECTED_SLIDE_COUNT:
        errors.append(f"Expected presentation to reference {EXPECTED_SLIDE_COUNT} slides, found {len(slide_ids)}")

    chart_targets: list[str] = []
    workbook_targets: list[str] = []
    for index, slide_id in enumerate(slide_ids, start=1):
        rel = presentation_rel_map.get(slide_id.get(f"{{{R_NS}}}id") or "")
        if rel is None or rel.get("Type") != SLIDE_REL:
            errors.append(f"Presentation slide {index} relationship is missing or has the wrong type")
            continue
        slide_part = _relationship_target(presentation_part, rel, parts, errors)
        if slide_part is None:
            continue
        slide_root = _parse_xml(parts, slide_part, errors)
        slide_rels_root = _parse_xml(parts, _rels_name(slide_part), errors)
        if slide_root is None or slide_rels_root is None:
            continue
        slide_rels = slide_rels_root.findall("rel:Relationship", REL)
        entry = ENTRY_BY_SLIDE.get(index)

        layout_rels = [r for r in slide_rels if r.get("Type") == SLIDE_LAYOUT_REL]
        if len(layout_rels) != 1:
            errors.append(f"Slide {index} must have exactly one layout relationship; found {len(layout_rels)}")
        elif entry is not None:
            actual_layout = _relationship_target(slide_part, layout_rels[0], parts, errors)
            actual_name = layout_name_by_part.get(actual_layout or "")
            if actual_name != entry.layout:
                errors.append(f"Slide {index} uses layout {actual_name!r} ({actual_layout}); expected {entry.layout!r}")

        _check_prohibited_copy(index, slide_root, errors)
        _check_no_hard_coded_style(f"Slide {index}", slide_root, errors)
        _check_placeholder_text_inherits(f"Slide {index}", slide_root, errors)

        image_rels = [r for r in slide_rels if r.get("Type") == IMAGE_REL]
        if image_rels or slide_root.findall(".//p:pic", NS):
            errors.append(f"Slide {index} contains a picture; example slides must not embed imagery")

        chart_rels = [r for r in slide_rels if r.get("Type") == CHART_REL]
        chart_refs = slide_root.findall(".//c:chart", NS)
        if index in NATIVE_CHART_SLIDES:
            if len(chart_rels) != 1 or len(chart_refs) != 1:
                errors.append(f"Native chart slide {index} must contain one chart relationship and one chart frame; "
                              f"found {len(chart_rels)} and {len(chart_refs)}")
            else:
                frame = chart_refs[0].getparent().getparent().getparent()
                if frame.find("p:nvGraphicFramePr/p:nvPr/p:ph", NS) is None:
                    errors.append(f"Native chart slide {index} chart is not bound to a content placeholder")
                chart_part = _relationship_target(slide_part, chart_rels[0], parts, errors)
                if chart_part:
                    chart_targets.append(chart_part)
                    workbook = _check_chart_part(index, chart_part, parts, overrides, defaults, errors)
                    if workbook:
                        workbook_targets.append(workbook)
        elif chart_rels or chart_refs:
            errors.append(f"Slide {index} is not a supported native-chart example but contains chart content")

        if index in TABLE_SLIDES:
            tables = slide_root.findall(".//a:tbl", NS)
            frames = [t.getparent().getparent().getparent() for t in tables]
            if len(tables) != 1 or frames[0].find("p:nvGraphicFramePr/p:nvPr/p:ph", NS) is None:
                errors.append(f"Table slide {index} must hold one table bound to a content placeholder")
            else:
                report["table_slide_count"] += 1
        if index in DIAGRAM_SLIDES:
            shapes = [s for s in slide_root.findall(".//p:sp", NS) + slide_root.findall(".//p:cxnSp", NS)
                      if s.find(".//p:nvPr/p:ph", NS) is None]
            if len(shapes) < MIN_DIAGRAM_SHAPES:
                errors.append(f"Diagram slide {index} has {len(shapes)} native shapes; expected at least "
                              f"{MIN_DIAGRAM_SHAPES}")
            else:
                report["diagram_slide_count"] += 1

    report["native_chart_part_count"] = len(chart_targets)
    report["embedded_workbook_count"] = len(workbook_targets)
    if len(chart_targets) != len(NATIVE_CHART_SLIDES) or len(set(chart_targets)) != len(chart_targets):
        errors.append(f"Expected {len(NATIVE_CHART_SLIDES)} distinct native chart parts, found {len(chart_targets)}")
    if len(workbook_targets) != len(NATIVE_CHART_SLIDES) or len(set(workbook_targets)) != len(workbook_targets):
        errors.append(f"Expected {len(NATIVE_CHART_SLIDES)} distinct embedded workbooks, found {len(workbook_targets)}")

    try:
        prs = Presentation(str(path))
        if len(prs.slide_masters) != 1:
            errors.append(f"python-pptx sees {len(prs.slide_masters)} masters; expected exactly one")
        if tuple(layout.name for layout in prs.slide_layouts) != expected_layout_names():
            errors.append("python-pptx layout names/order do not match the catalogue")
        if len(prs.slides) != EXPECTED_SLIDE_COUNT:
            errors.append(f"python-pptx sees {len(prs.slides)} slides; expected exactly {EXPECTED_SLIDE_COUNT}")
    except Exception as exc:  # pragma: no cover - exact package errors vary by dependency version.
        errors.append(f"python-pptx could not load the deck: {type(exc).__name__}: {exc}")

    report["ok"] = not errors
    return report


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pptx", nargs="?", type=Path, default=DEFAULT_DECK, help="Generated fallback deck to verify")
    parser.add_argument("--warn-only", action="store_true", help="Print failures but exit successfully")
    args = parser.parse_args(list(argv) if argv is not None else None)
    report = inspect_deck(args.pptx)
    print(json.dumps(report, indent=2))
    return 0 if report["ok"] or args.warn_only else 1


if __name__ == "__main__":
    raise SystemExit(main())
