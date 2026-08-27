from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER


REPO_ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = REPO_ROOT / "scripts" / "inspect_template.py"
SPEC = importlib.util.spec_from_file_location("inspect_template", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
inspect_template = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(inspect_template)


def make_shape(
    placeholder_type: PP_PLACEHOLDER,
    *,
    name: str = "Text Placeholder 1",
    has_text_frame: bool = True,
    idx: int = 1,
) -> SimpleNamespace:
    return SimpleNamespace(
        placeholder_format=SimpleNamespace(idx=idx, type=placeholder_type),
        name=name,
        has_text_frame=has_text_frame,
        left=914400,
        top=1828800,
        width=2743200,
        height=3657600,
    )


class PlaceholderClassificationTests(unittest.TestCase):
    def test_definitive_placeholder_types(self) -> None:
        cases = (
            (PP_PLACEHOLDER.TITLE, "title", "core_content", True),
            (PP_PLACEHOLDER.SUBTITLE, "subtitle", "optional_supporting", True),
            (PP_PLACEHOLDER.FOOTER, "footer_or_admin", "administrative", True),
            (PP_PLACEHOLDER.CHART, "non_text", "non_text", False),
        )
        for placeholder_type, role, content_class, is_text in cases:
            with self.subTest(placeholder_type=placeholder_type):
                result = inspect_template.classify_placeholder(
                    make_shape(placeholder_type, has_text_frame=is_text)
                )
                self.assertEqual(result["content_role"], role)
                self.assertEqual(result["content_class"], content_class)
                self.assertEqual(result["is_text_placeholder"], is_text)
                self.assertFalse(result["review_required"])

    def test_specific_names_map_to_optional_supporting_roles(self) -> None:
        cases = (
            ("Eyebrow Text", "section_label", "optional_supporting"),
            ("Image Caption", "caption", "optional_supporting"),
            ("Source Note", "source", "optional_supporting"),
            ("Quote Attribution", "attribution", "optional_supporting"),
            ("Customer Quotation", "quote", "optional_supporting"),
            ("Key Takeaway", "callout", "optional_supporting"),
            ("Legal Footer", "footer_or_admin", "administrative"),
        )
        for name, role, content_class in cases:
            with self.subTest(name=name):
                result = inspect_template.classify_placeholder(
                    make_shape(PP_PLACEHOLDER.BODY, name=name)
                )
                self.assertEqual(result["content_role"], role)
                self.assertEqual(result["content_class"], content_class)
                self.assertEqual(result["role_basis"], "placeholder_name")
                self.assertTrue(result["review_required"])

    def test_generic_body_is_core_destination_not_a_fill_requirement(self) -> None:
        result = inspect_template.classify_placeholder(
            make_shape(PP_PLACEHOLDER.BODY)
        )
        self.assertEqual(result["content_role"], "body_content")
        self.assertEqual(result["content_class"], "core_content")
        self.assertFalse(result["review_required"])

    def test_multipurpose_text_placeholder_requires_review(self) -> None:
        result = inspect_template.classify_placeholder(
            make_shape(PP_PLACEHOLDER.OBJECT, name="Content Placeholder 2")
        )
        self.assertEqual(result["content_role"], "body_content")
        self.assertEqual(result["content_class"], "review")
        self.assertEqual(result["role_confidence"], "low")
        self.assertTrue(result["review_required"])

    def test_name_matching_uses_word_boundaries(self) -> None:
        result = inspect_template.classify_placeholder(
            make_shape(PP_PLACEHOLDER.BODY, name="Highlighted content")
        )
        self.assertEqual(result["content_role"], "body_content")


class RoleOverrideTests(unittest.TestCase):
    def write_override(self, directory: Path, data: object) -> Path:
        path = directory / "roles.json"
        path.write_text(json.dumps(data), encoding="utf-8")
        return path

    def test_valid_override_is_loaded_and_applied(self) -> None:
        data = {
            "version": 1,
            "overrides": {
                "2:7": {
                    "layout_name": "Picture with Caption",
                    "placeholder_name": "Text Placeholder 4",
                    "content_role": "caption",
                    "content_class": "optional_supporting",
                    "note": "Use only with supplied caption text.",
                }
            },
        }
        with tempfile.TemporaryDirectory() as temp_directory:
            path = self.write_override(Path(temp_directory), data)
            overrides = inspect_template.load_role_overrides(path)

        result = inspect_template.placeholder_info(
            make_shape(
                PP_PLACEHOLDER.BODY,
                name="Text Placeholder 4",
                idx=7,
            ),
            layout_index=2,
            layout_name="Picture with Caption",
            override=overrides[(2, 7)],
        )
        self.assertEqual(result["content_role"], "caption")
        self.assertEqual(result["role_basis"], "override")
        self.assertEqual(result["role_confidence"], "explicit")
        self.assertFalse(result["review_required"])
        self.assertEqual(result["role_note"], "Use only with supplied caption text.")

    def test_invalid_schema_role_class_and_field_are_rejected(self) -> None:
        cases = (
            {"version": 2, "overrides": {}},
            {
                "version": 1,
                "overrides": {
                    "0:1": {
                        "content_role": "strapline",
                        "content_class": "optional_supporting",
                    }
                },
            },
            {
                "version": 1,
                "overrides": {
                    "0:1": {
                        "content_role": "caption",
                        "content_class": "required",
                    }
                },
            },
            {"version": 1, "overrides": {}, "unexpected": True},
        )
        with tempfile.TemporaryDirectory() as temp_directory:
            directory = Path(temp_directory)
            for index, data in enumerate(cases):
                with self.subTest(index=index):
                    path = directory / f"roles-{index}.json"
                    path.write_text(json.dumps(data), encoding="utf-8")
                    with self.assertRaises(ValueError):
                        inspect_template.load_role_overrides(path)

    def test_duplicate_json_keys_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_directory:
            path = Path(temp_directory) / "roles.json"
            path.write_text(
                '{"version": 1, "version": 1, "overrides": {}}',
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "Duplicate JSON key"):
                inspect_template.load_role_overrides(path)

    def test_guard_name_mismatch_is_rejected(self) -> None:
        override = {
            "layout_name": "Expected Layout",
            "placeholder_name": "Expected Placeholder",
            "content_role": "caption",
            "content_class": "optional_supporting",
        }
        with self.assertRaisesRegex(ValueError, "expected layout"):
            inspect_template.placeholder_info(
                make_shape(PP_PLACEHOLDER.BODY, name="Expected Placeholder"),
                layout_index=0,
                layout_name="Changed Layout",
                override=override,
            )

    def test_missing_override_target_is_rejected(self) -> None:
        override = {
            "content_role": "caption",
            "content_class": "optional_supporting",
        }
        with tempfile.TemporaryDirectory() as temp_directory:
            template_path = Path(temp_directory) / "template.pptx"
            Presentation().save(template_path)
            with self.assertRaisesRegex(ValueError, "not found"):
                inspect_template.inspect(template_path, {(999, 999): override})


class OutputTests(unittest.TestCase):
    def test_markdown_includes_reviewable_role_fields(self) -> None:
        report = {
            "path": "template.pptx",
            "slide_width_in": 13.333,
            "slide_height_in": 7.5,
            "layout_count": 1,
            "role_override_count": 0,
            "review_required_count": 1,
            "layouts": [
                {
                    "index": 0,
                    "name": "Title | Content",
                    "placeholder_count": 1,
                    "review_required_count": 1,
                    "placeholders": [
                        {
                            "idx": 1,
                            "type": "OBJECT (7)",
                            "name": "Content | Placeholder",
                            "content_role": "body_content",
                            "content_class": "review",
                            "role_basis": "ambiguous_placeholder_type",
                            "review_required": True,
                            "left_in": 1.0,
                            "top_in": 2.0,
                            "width_in": 3.0,
                            "height_in": 4.0,
                        }
                    ],
                }
            ],
        }
        markdown = inspect_template.to_markdown(report)
        self.assertIn("Placeholders requiring role review: 1", markdown)
        self.assertIn("| role | class | basis | review |", markdown)
        self.assertIn("body_content", markdown)
        self.assertIn("Content \\| Placeholder", markdown)

    def test_inspect_reports_role_counts_and_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temp_directory:
            template_path = Path(temp_directory) / "template.pptx"
            Presentation().save(template_path)
            report = inspect_template.inspect(template_path)

        self.assertIn("review_required_count", report)
        self.assertIn("role_override_count", report)
        placeholders = [
            placeholder
            for layout in report["layouts"]
            for placeholder in layout["placeholders"]
        ]
        self.assertTrue(placeholders)
        for placeholder in placeholders:
            self.assertIn("content_role", placeholder)
            self.assertIn("content_class", placeholder)
            self.assertIn("role_basis", placeholder)
            self.assertIn("role_confidence", placeholder)
            self.assertIn("review_required", placeholder)


if __name__ == "__main__":
    unittest.main()
